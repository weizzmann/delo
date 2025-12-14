from flask import Blueprint, send_file, current_app, request
from flask_login import login_required, current_user
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from ..models import Case, Volume, VolumeDocument
from ..extensions import db

api_export = Blueprint('api_export', __name__)


@api_export.route('/cases/<int:case_id>/export-excel', methods=['GET'])
@login_required
def export_case(case_id):
    case = Case.query.get_or_404(case_id)
    if case.department_id != current_user.department_id:
        return {'error': 'forbidden'}, 403
    
    # Получаем номер тома из параметра (если указан)
    volume_number = request.args.get('volume', type=int)
    
    wb = Workbook()
    # Удаляем стандартный лист
    wb.remove(wb.active)
    
    # Стили для форматирования
    header_font = Font(bold=True, size=12)
    cell_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    # Если указан конкретный том
    if volume_number:
        volume = Volume.query.filter_by(case_id=case_id, number=volume_number).first()
        if not volume:
            return {'error': 'volume not found'}, 404
        
        ws = wb.create_sheet(title=f'Том_{volume.number}')
        
        # Заголовок
        ws.merge_cells('A1:F1')
        ws['A1'] = f'Дело: {case.name} - Том {volume.number}'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('A2:F2')
        ws['A2'] = f'Описание: {case.description or "Нет описания"}'
        ws['A2'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('A3:F3')
        ws['A3'] = f'Страницы в томе: {volume.pages_start}-{volume.pages_end} (всего: {volume.total_pages} стр.)'
        ws['A3'].alignment = Alignment(horizontal='center')
        
        # Пустая строка
        ws.append([])
        
        # Заголовки таблицы
        headers = ['№ в томе', 'Название документа', 'Входящий №', 'Исходящий №', 'Доп. информация', 'Страницы в томе']
        ws.append(headers)
        
        # Форматирование заголовков
        for col in range(1, 7):
            cell = ws.cell(row=5, column=col)
            cell.font = header_font
            cell.border = cell_border
            cell.alignment = center_alignment
            ws.column_dimensions[chr(64 + col)].width = 20
        
        # Данные документов
        row_num = 6
        for vd in volume.documents:
            ws.append([
                vd.order_in_volume,
                vd.document.title,
                vd.document.incoming_number or '',
                vd.document.outgoing_number or '',
                vd.document.extra_info or '',
                f"{vd.start_page}-{vd.end_page}",
            ])
            
            # Форматирование строки
            for col in range(1, 7):
                cell = ws.cell(row=row_num, column=col)
                cell.border = cell_border
                if col == 1:  # Номер в томе
                    cell.alignment = center_alignment
            row_num += 1
        
        # Итоговая строка
        ws.append([])
        ws.merge_cells(f'A{row_num}:F{row_num}')
        ws[f'A{row_num}'] = f'Итого документов: {len(volume.documents)}, Всего страниц: {volume.total_pages}'
        ws[f'A{row_num}'].font = Font(bold=True)
        ws[f'A{row_num}'].alignment = Alignment(horizontal='center')
        
    else:
        # Экспорт всех томов
        for volume in sorted(case.volumes, key=lambda v: v.number):
            ws = wb.create_sheet(title=f'Том_{volume.number}')
            
            # Заголовок
            ws.merge_cells('A1:F1')
            ws['A1'] = f'Дело: {case.name} - Том {volume.number}'
            ws['A1'].font = Font(bold=True, size=14)
            ws['A1'].alignment = Alignment(horizontal='center')
            
            ws.merge_cells('A2:F2')
            ws['A2'] = f'Описание: {case.description or "Нет описания"}'
            ws['A2'].alignment = Alignment(horizontal='center')
            
            ws.merge_cells('A3:F3')
            ws['A3'] = f'Страницы в томе: {volume.pages_start}-{volume.pages_end} (всего: {volume.total_pages} стр.)'
            ws['A3'].alignment = Alignment(horizontal='center')
            
            # Пустая строка
            ws.append([])
            
            # Заголовки таблицы
            headers = ['№ в томе', 'Название документа', 'Входящий №', 'Исходящий №', 'Доп. информация', 'Страницы в томе']
            ws.append(headers)
            
            # Форматирование заголовков
            for col in range(1, 7):
                cell = ws.cell(row=5, column=col)
                cell.font = header_font
                cell.border = cell_border
                cell.alignment = center_alignment
                ws.column_dimensions[chr(64 + col)].width = 20
            
            # Данные документов
            row_num = 6
            for vd in volume.documents:
                ws.append([
                    vd.order_in_volume,
                    vd.document.title,
                    vd.document.incoming_number or '',
                    vd.document.outgoing_number or '',
                    vd.document.extra_info or '',
                    f"{vd.start_page}-{vd.end_page}",
                ])
                
                # Форматирование строки
                for col in range(1, 7):
                    cell = ws.cell(row=row_num, column=col)
                    cell.border = cell_border
                    if col == 1:  # Номер в томе
                        cell.alignment = center_alignment
                row_num += 1
            
            # Итоговая строка
            ws.append([])
            ws.merge_cells(f'A{row_num}:F{row_num}')
            ws[f'A{row_num}'] = f'Итого документов: {len(volume.documents)}, Всего страниц: {volume.total_pages}'
            ws[f'A{row_num}'].font = Font(bold=True)
            ws[f'A{row_num}'].alignment = Alignment(horizontal='center')
        
        # Добавляем сводный лист
        summary_ws = wb.create_sheet(title='Сводка', index=0)
        
        summary_ws.merge_cells('A1:E1')
        summary_ws['A1'] = f'Сводка по делу: {case.name}'
        summary_ws['A1'].font = Font(bold=True, size=16)
        summary_ws['A1'].alignment = Alignment(horizontal='center')
        
        summary_ws.merge_cells('A2:E2')
        summary_ws['A2'] = f'Описание: {case.description or "Нет описания"}'
        summary_ws['A2'].alignment = Alignment(horizontal='center')
        
        summary_ws.merge_cells('A3:E3')
        summary_ws['A3'] = f'Департамент: {case.department.name}'
        summary_ws['A3'].alignment = Alignment(horizontal='center')
        
        summary_ws.append([])
        
        # Заголовки сводной таблицы
        summary_headers = ['Том', 'Статус', 'Страниц в томе', 'Документов', 'Заполнение']
        summary_ws.append(summary_headers)
        
        for col in range(1, 6):
            cell = summary_ws.cell(row=6, column=col)
            cell.font = header_font
            cell.border = cell_border
            cell.alignment = center_alignment
            summary_ws.column_dimensions[chr(64 + col)].width = 20
        
        # Данные по томам
        row_num = 7
        total_documents = 0
        total_pages = 0
        
        for volume in sorted(case.volumes, key=lambda v: v.number):
            documents_count = len(volume.documents)
            total_documents += documents_count
            total_pages += volume.total_pages
            
            # Определяем статус
            if volume.total_pages == 250:
                status = 'Полный'
                fill_percent = '100%'
            else:
                status = 'Частично заполнен'
                fill_percent = f'{int((volume.total_pages / 250) * 100)}%'
            
            summary_ws.append([
                f'Том {volume.number}',
                status,
                volume.total_pages,
                documents_count,
                fill_percent
            ])
            
            for col in range(1, 6):
                cell = summary_ws.cell(row=row_num, column=col)
                cell.border = cell_border
                if col in [1, 3, 4]:
                    cell.alignment = center_alignment
            row_num += 1
        
        # Итоговая строка
        summary_ws.append([])
        summary_ws.merge_cells(f'A{row_num}:E{row_num}')
        summary_ws[f'A{row_num}'] = f'ИТОГО: {len(case.volumes)} томов, {total_documents} документов, {total_pages} страниц'
        summary_ws[f'A{row_num}'].font = Font(bold=True, size=12)
        summary_ws[f'A{row_num}'].alignment = Alignment(horizontal='center')
    
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    
    if volume_number:
        filename = f"Дело_{case.name}_Том_{volume_number}.xlsx"
    else:
        filename = f"Дело_{case.name}_Все_тома.xlsx"
    
    # Заменяем недопустимые символы в имени файла
    filename = filename.replace('/', '_').replace('\\', '_').replace(':', '_')
    
    return send_file(
        bio, 
        download_name=filename, 
        as_attachment=True, 
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )